"""Bounded, read-only local CSV, SQLite, and Excel analysis helpers."""
from __future__ import annotations

import csv
import json
import re
import sqlite3
import threading
import zipfile
from pathlib import Path
from typing import Any

from config import BASE_DIR, WORKSPACE_DIR

MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_ROWS = 1000
_FORBIDDEN_SQL = re.compile(r"\b(attach|detach|pragma|copy|install|load|call|create|drop|alter|insert|update|delete|replace|vacuum|execute|set|reset|extension|read_csv|read_json|read_parquet|glob)\b", re.I)


def _allowed_file(raw_path: str, suffixes: set[str]) -> Path:
    candidate = Path(raw_path).expanduser().resolve(strict=True)
    roots = [Path(BASE_DIR).resolve(), Path(WORKSPACE_DIR).resolve()]
    if not any(candidate == root or root in candidate.parents for root in roots):
        raise ValueError("File must be inside the project folder or configured workspace.")
    if not candidate.is_file() or candidate.suffix.lower() not in suffixes:
        raise ValueError(f"Expected a regular file with extension in {sorted(suffixes)}.")
    if candidate.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(f"File exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MiB analysis limit.")
    return candidate


def _validate_select(sql_query: str) -> str:
    sql = (sql_query or "").strip()
    if not sql or len(sql) > 8000 or sql.count(";") > 1 or (";" in sql and not sql.endswith(";")):
        raise ValueError("Provide one SELECT query of at most 8,000 characters.")
    sql = sql.rstrip(";").strip()
    if not re.match(r"^(select|with)\b", sql, re.I) or _FORBIDDEN_SQL.search(sql):
        raise ValueError("Only read-only SELECT queries are allowed; DDL, mutations, and file/network functions are blocked.")
    return sql


def query_csv(file_path: str, sql_query: str) -> str:
    """Run a read-only DuckDB query against one CSV registered as table `data`."""
    path = _allowed_file(file_path, {".csv", ".tsv"})
    sql = _validate_select(sql_query)
    try:
        import duckdb
    except ImportError:
        return "DuckDB is not installed. Install the project requirement `duckdb` to use CSV SQL queries."
    try:
        connection = duckdb.connect(":memory:", config={"threads": "2"})
        def interrupt_query() -> None:
            try:
                connection.interrupt()
            except Exception:
                pass

        timer = threading.Timer(20, interrupt_query)
        timer.daemon = True
        timer.start()
        try:
            delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
            # Use the Python API for this one approved source, then disable file access.
            connection.read_csv(str(path), delimiter=delimiter, header=True,
                                sample_size=10000).create("data")
            connection.execute("SET enable_external_access = false")
            cursor = connection.execute(sql)
            names = [item[0] for item in (cursor.description or [])]
            rows = cursor.fetchmany(MAX_ROWS + 1)
            truncated = len(rows) > MAX_ROWS
            rows = rows[:MAX_ROWS]
            result = {"columns": names, "rows": [list(row) for row in rows], "truncated": truncated}
            return json.dumps(result, ensure_ascii=False, default=str)
        finally:
            timer.cancel()
            connection.close()
    except Exception as exc:
        return f"CSV query failed: {type(exc).__name__}: {exc}"


def query_sqlite_db(db_path: str, sql_query: str) -> str:
    """Execute a bounded read-only SELECT against an approved local SQLite file."""
    path = _allowed_file(db_path, {".sqlite", ".sqlite3", ".db"})
    sql = _validate_select(sql_query)
    conn = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True, timeout=3)
    conn.row_factory = sqlite3.Row
    conn.set_progress_handler(lambda: 1, 2_000_000)
    try:
        cur = conn.execute(sql)
        rows = cur.fetchmany(MAX_ROWS + 1)
        truncated = len(rows) > MAX_ROWS
        output = {"columns": [d[0] for d in cur.description or []],
                  "rows": [dict(row) for row in rows[:MAX_ROWS]], "truncated": truncated}
        return json.dumps(output, ensure_ascii=False, default=str)
    except Exception as exc:
        return f"SQLite query failed: {type(exc).__name__}: {exc}"
    finally:
        conn.close()


def analyze_excel(file_path: str, operation: str = "preview", sheet_name: str = "") -> str:
    """Read an Excel sheet without evaluating formulas; supports preview/profile/list_sheets."""
    path = _allowed_file(file_path, {".xlsx", ".xlsm"})
    if operation not in {"preview", "profile", "list_sheets"}:
        return "operation must be preview, profile, or list_sheets."
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            expanded = sum(item.file_size for item in members)
            if expanded > 250 * 1024 * 1024 or any(
                item.file_size > 0 and item.compress_size == 0
                or item.compress_size > 0 and item.file_size / item.compress_size > 200
                for item in members
            ):
                return "Excel archive rejected: expanded size or compression ratio exceeds the safety limit."
    except (OSError, zipfile.BadZipFile):
        return "Excel archive is invalid or unreadable."
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "openpyxl is not installed. Install the project requirement `openpyxl` to analyze Excel files."
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        if operation == "list_sheets":
            return json.dumps({"sheets": workbook.sheetnames})
        if sheet_name and sheet_name not in workbook.sheetnames:
            return f"Unknown sheet. Available sheets: {', '.join(workbook.sheetnames)}"
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(v) if v is not None else f"column_{i + 1}" for i, v in enumerate(next(iterator, ()))]
        rows = []
        count = 0
        numeric: dict[str, list[float]] = {h: [] for h in headers}
        for row in iterator:
            if all(v is None for v in row):
                continue
            count += 1
            if len(rows) < MAX_ROWS:
                rows.append(dict(zip(headers, row)))
            for h, value in zip(headers, row):
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    numeric[h].append(float(value))
            if count >= 100_000:
                break
        if operation == "preview":
            result: dict[str, Any] = {"sheet": sheet.title, "columns": headers, "rows": rows, "truncated": count > MAX_ROWS}
        else:
            result = {"sheet": sheet.title, "rows_scanned": count, "numeric_columns": {
                h: {"count": len(vals), "min": min(vals), "max": max(vals), "mean": sum(vals) / len(vals)}
                for h, vals in numeric.items() if vals}}
        workbook.close()
        return json.dumps(result, ensure_ascii=False, default=str)
    except Exception as exc:
        return f"Excel analysis failed: {type(exc).__name__}: {exc}"
